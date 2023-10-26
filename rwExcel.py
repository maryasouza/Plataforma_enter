from openpyxl.workbook import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


class Encrypt:
    def __init__(self, senha):
        self.senha = senha


class ExcelExporter:
    def __init__(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Senha Encrypt"

    def write_header_row(self):
        header = ["senha"]
        self.ws.append(header)

    def write_senha_data(self, senha_data):
        self.ws.append([senha_data])

    def add_table(self):
        table = Table(displayName="AccountTable", ref=self.ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False)
        self.ws.add_table(table)

    def save_workbook(self, output_file):
        self.wb.save(output_file)


class SenhaInformationExporter:
    def __init__(self, excel_exporter, senha_out):
        self.excel_exporter = excel_exporter
        self.senha_out = senha_out

    def export_senha_information(self, output_file):
        senha_data = self.senha_out
        self.excel_exporter.write_header_row()
        self.excel_exporter.write_senha_data(senha_data)
        self.excel_exporter.add_table()
        self.excel_exporter.save_workbook(output_file)

